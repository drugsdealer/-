"""
Импорт данных из Excel-файлов в PostgreSQL.
Запуск: python import_data.py
"""
import os
import sys
import hashlib
import psycopg2
import openpyxl

# ─── Настройки подключения ────────────────────────────────────────────────────
DB_CONFIG = {
    "host":     os.getenv("DB_HOST",     "localhost"),
    "port":     int(os.getenv("DB_PORT", "5432")),
    "dbname":   os.getenv("DB_NAME",     "stroy_materials"),
    "user":     os.getenv("DB_USER",     "postgres"),
    "password": os.getenv("DB_PASSWORD", "postgres"),
}

IMPORT_DIR = os.path.join(os.path.dirname(__file__), "import")

PRODUCTS_FILE      = os.path.join(IMPORT_DIR, "Tovar.xlsx")
USERS_FILE         = os.path.join(IMPORT_DIR, "user_import.xlsx")
ORDERS_FILE        = os.path.join(IMPORT_DIR, "Заказ_import.xlsx")
PICKUP_POINTS_FILE = os.path.join(IMPORT_DIR, "Пункты выдачи_import.xlsx")


def hash_password(pwd: str) -> str:
    return hashlib.sha256(pwd.encode("utf-8")).hexdigest()


def get_or_create(cursor, table: str, id_col: str, name_col: str, name: str) -> int:
    cursor.execute(f"SELECT {id_col} FROM {table} WHERE {name_col} = %s", (name,))
    row = cursor.fetchone()
    if row:
        return row[0]
    cursor.execute(
        f"INSERT INTO {table} ({name_col}) VALUES (%s) RETURNING {id_col}",
        (name,),
    )
    return cursor.fetchone()[0]


def import_products(cursor):
    if not os.path.exists(PRODUCTS_FILE):
        print(f"  ФАЙЛ НЕ НАЙДЕН: {PRODUCTS_FILE}")
        return

    wb = openpyxl.load_workbook(PRODUCTS_FILE, data_only=True)
    ws = wb.active
    count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        article, name, unit, price, supplier, manufacturer, category, discount, qty, description, image = row

        if not article or not name:
            continue

        unit_id         = get_or_create(cursor, "units",         "unit_id",         "unit_name",         str(unit or "шт."))
        supplier_id     = get_or_create(cursor, "suppliers",     "supplier_id",     "supplier_name",     str(supplier or "Не указан"))
        manufacturer_id = get_or_create(cursor, "manufacturers", "manufacturer_id", "manufacturer_name", str(manufacturer or "Не указан"))
        category_id     = get_or_create(cursor, "categories",    "category_id",     "category_name",     str(category or "Прочее"))

        cursor.execute("""
            INSERT INTO products
                (article, product_name, unit_id, price, supplier_id, manufacturer_id,
                 category_id, discount, quantity_stock, description, image_path)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            ON CONFLICT (article) DO UPDATE SET
                product_name     = EXCLUDED.product_name,
                price            = EXCLUDED.price,
                discount         = EXCLUDED.discount,
                quantity_stock   = EXCLUDED.quantity_stock,
                description      = EXCLUDED.description,
                image_path       = EXCLUDED.image_path
        """, (
            str(article).strip(),
            str(name).strip(),
            unit_id,
            float(price or 0),
            supplier_id,
            manufacturer_id,
            category_id,
            float(discount or 0),
            int(qty or 0),
            str(description).strip() if description else None,
            str(image).strip() if image else None,
        ))
        count += 1

    print(f"  Товары: импортировано {count} записей")


def import_users(cursor):
    if not os.path.exists(USERS_FILE):
        print(f"  ФАЙЛ НЕ НАЙДЕН: {USERS_FILE}")
        return

    wb = openpyxl.load_workbook(USERS_FILE, data_only=True)
    ws = wb.active
    count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        role_name, full_name, login, password = row

        if not login or not full_name:
            continue

        cursor.execute("SELECT role_id FROM roles WHERE role_name = %s", (str(role_name).strip(),))
        role_row = cursor.fetchone()
        if not role_row:
            print(f"    ВНИМАНИЕ: роль '{role_name}' не найдена, пропуск пользователя {login}")
            continue
        role_id = role_row[0]

        cursor.execute("""
            INSERT INTO users (full_name, login, password, role_id)
            VALUES (%s, %s, %s, %s)
            ON CONFLICT (login) DO UPDATE SET
                full_name = EXCLUDED.full_name,
                role_id   = EXCLUDED.role_id
        """, (
            str(full_name).strip(),
            str(login).strip(),
            hash_password(str(password).strip()),
            role_id,
        ))
        count += 1

    print(f"  Пользователи: импортировано {count} записей")


def import_pickup_points(cursor):
    if not os.path.exists(PICKUP_POINTS_FILE):
        print(f"  ФАЙЛ НЕ НАЙДЕН: {PICKUP_POINTS_FILE}")
        print("  Создаю тестовые пункты выдачи...")
        sample_points = [
            "г. Москва, ул. Строителей, д. 1",
            "г. Москва, пр. Ленина, д. 25",
            "г. Москва, ул. Гагарина, д. 10",
            "г. Москва, ул. Мира, д. 5",
            "г. Москва, ул. Победы, д. 3",
        ]
        for address in sample_points:
            cursor.execute(
                "INSERT INTO pickup_points (address) VALUES (%s) ON CONFLICT (address) DO NOTHING",
                (address,),
            )
        print(f"  Пункты выдачи: создано {len(sample_points)} записей")
        return

    wb = openpyxl.load_workbook(PICKUP_POINTS_FILE, data_only=True)
    ws = wb.active
    count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        address = str(row[0]).strip()
        cursor.execute(
            "INSERT INTO pickup_points (address) VALUES (%s) ON CONFLICT (address) DO NOTHING",
            (address,),
        )
        count += 1

    print(f"  Пункты выдачи: импортировано {count} записей")


def import_orders(cursor):
    if not os.path.exists(ORDERS_FILE):
        print(f"  ФАЙЛ НЕ НАЙДЕН: {ORDERS_FILE}")
        print("  Создаю тестовые заказы...")

        cursor.execute("SELECT user_id FROM users LIMIT 5")
        user_ids = [r[0] for r in cursor.fetchall()]

        cursor.execute("SELECT pickup_point_id FROM pickup_points LIMIT 3")
        point_ids = [r[0] for r in cursor.fetchall()]

        cursor.execute("SELECT status_id FROM order_statuses WHERE status_name = 'Новый'")
        status_id = cursor.fetchone()[0]

        cursor.execute("SELECT product_id, price FROM products LIMIT 10")
        prods = cursor.fetchall()

        if not user_ids or not point_ids or not prods:
            print("  Нет данных для создания заказов")
            return

        import datetime, random
        for i in range(5):
            user_id = user_ids[i % len(user_ids)]
            point_id = point_ids[i % len(point_ids)]
            order_date = datetime.datetime.now() - datetime.timedelta(days=i * 3)
            cursor.execute("""
                INSERT INTO orders (user_id, order_date, pickup_point_id, status_id, total_amount)
                VALUES (%s, %s, %s, %s, 0) RETURNING order_id
            """, (user_id, order_date, point_id, status_id))
            order_id = cursor.fetchone()[0]

            total = 0
            for j in range(1, 3):
                prod_id, prod_price = prods[(i + j) % len(prods)]
                qty = j + 1
                cursor.execute("""
                    INSERT INTO order_items (order_id, product_id, quantity, price)
                    VALUES (%s, %s, %s, %s)
                """, (order_id, prod_id, qty, prod_price))
                total += prod_price * qty

            cursor.execute("UPDATE orders SET total_amount = %s WHERE order_id = %s", (total, order_id))

        print("  Заказы: создано 5 тестовых записей")
        return

    wb = openpyxl.load_workbook(ORDERS_FILE, data_only=True)
    ws = wb.active
    count = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        # Ожидаемые столбцы: user_login, order_date, pickup_address, status, product_article, quantity
        user_login, order_date, pickup_address, status_name, article, qty = (row + (None,) * 6)[:6]

        cursor.execute("SELECT user_id FROM users WHERE login = %s", (str(user_login or "").strip(),))
        user_row = cursor.fetchone()
        if not user_row:
            continue

        pickup_id = None
        if pickup_address:
            cursor.execute("SELECT pickup_point_id FROM pickup_points WHERE address = %s", (str(pickup_address).strip(),))
            pp = cursor.fetchone()
            if pp:
                pickup_id = pp[0]

        cursor.execute("SELECT status_id FROM order_statuses WHERE status_name = %s", (str(status_name or "Новый").strip(),))
        st = cursor.fetchone()
        status_id = st[0] if st else 1

        cursor.execute("""
            INSERT INTO orders (user_id, order_date, pickup_point_id, status_id, total_amount)
            VALUES (%s, %s, %s, %s, 0) RETURNING order_id
        """, (user_row[0], order_date, pickup_id, status_id))
        order_id = cursor.fetchone()[0]

        if article:
            cursor.execute("SELECT product_id, price FROM products WHERE article = %s", (str(article).strip(),))
            prod = cursor.fetchone()
            if prod:
                cursor.execute(
                    "INSERT INTO order_items (order_id, product_id, quantity, price) VALUES (%s, %s, %s, %s)",
                    (order_id, prod[0], int(qty or 1), prod[1]),
                )
                cursor.execute(
                    "UPDATE orders SET total_amount = %s WHERE order_id = %s",
                    (prod[1] * int(qty or 1), order_id),
                )
        count += 1

    print(f"  Заказы: импортировано {count} записей")


def main():
    print("Подключение к PostgreSQL...")
    try:
        conn = psycopg2.connect(**DB_CONFIG)
        conn.autocommit = False
        cursor = conn.cursor()
        print("Подключено.\n")
    except Exception as exc:
        print(f"ОШИБКА подключения: {exc}")
        sys.exit(1)

    try:
        print("Импорт данных:")
        import_products(cursor)
        import_users(cursor)
        import_pickup_points(cursor)
        import_orders(cursor)

        conn.commit()
        print("\nИмпорт завершён успешно.")
    except Exception as exc:
        conn.rollback()
        print(f"\nОШИБКА: {exc}")
        raise
    finally:
        cursor.close()
        conn.close()


if __name__ == "__main__":
    main()
